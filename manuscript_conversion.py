"""在 Markdown、Nature Article 投稿用 Word 和表格文件之间保留可编辑结构。"""

from __future__ import annotations

import base64
import copy
import errno
import html
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zlib
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from 转换工具.ppt_figures import export_ppt_figures, picture_markup, read_ppt_figures


ROOT = Path(__file__).resolve().parent
DEPENDENCIES = ROOT / ".manuscript_deps"
TEMP_ROOT = ROOT / "temp"
CSL = ROOT / "nature_article.csl"
OPERATION_REFERENCE = Path(r"D:\OneDrive\Operation.docx")
CITATION_STYLE_PREFIX = "MDCITE_"
FIGURE_REFERENCE_STYLE_PREFIX = "MDFIGREF_"
BIBLIOGRAPHY_STYLE = "MDConverterBibliography"
PANDOC_INPUT = "markdown+pipe_tables+raw_html+tex_math_dollars+citations"
PANDOC_OUTPUT = "markdown+pipe_tables+raw_html-simple_tables-grid_tables-multiline_tables"
XLSX_DIRECTIVE = re.compile(r"<!--\s*xlsx-table\s+(\{.*?\})\s*-->", re.I | re.S)
HTML_TABLE = re.compile(r"<table\b[^>]*>.*?</table\s*>", re.I | re.S)
MARKED_IMAGE = re.compile(r"(!\[[^\]]*\]\()&&\"([^\"]+)\"(\))")
STANDALONE_IMAGE = re.compile(r"(?m)^\s*&&\"([^\"]+)\"\s*$")
PPT_FIGURE = re.compile(r'(?m)^[ \t]*\{\{pptfig:"([^"\r\n]+)"\|([^|}\r\n]+)\}\}[ \t]*$')
PPT_TITLE_PREFIX = "MDPPTFIG_"
PPT_EDITED_TITLE_PREFIX = "MDPPTEDIT_"
PPT_EDITED_COMMENT = re.compile(r"<!--\s*pptfig-edited\s+(\{.*\})\s*-->", re.S)
PPT_EDITED_SOURCE = re.compile(r"<!--\s*pptfig-edited\s+(\{.*?\})\s*-->", re.S)
FIGURE_REFERENCE = re.compile(r"\{\{figref:([^\s|{}]+)(?:\|([a-z](?:[,–-][a-z])*))?\}\}")
MANUAL_FIGURE_NUMBER = re.compile(r"^(?:图\s*\d+|Fig(?:ure)?\.?\s*\d+)\s*[｜|:：]\s*", re.I)


def ensure_dependencies() -> str:
    """准备转换所需依赖，并返回内置 Pandoc 可执行文件路径。"""
    pandoc = DEPENDENCIES / "pypandoc" / "files" / "pandoc.exe"
    required = ("docx", "openpyxl", "lxml", "pptx", "PIL")
    if pandoc.is_file() and all(importlib.util.find_spec(name) for name in required):
        return str(pandoc)
    if DEPENDENCIES.is_dir() and str(DEPENDENCIES) not in sys.path:
        sys.path.append(str(DEPENDENCIES))
    if pandoc.is_file() and all(importlib.util.find_spec(name) for name in required):
        return str(pandoc)
    if not pandoc.is_file() or any(importlib.util.find_spec(name) is None for name in required):
        print("首次运行：正在安装转换依赖到脚本旁的 .manuscript_deps 目录……")
        command = [
            sys.executable, "-m", "pip", "install", "--disable-pip-version-check",
            "--target", str(DEPENDENCIES), "pypandoc_binary==1.17",
            "python-docx==1.2.0", "openpyxl==3.1.5", "lxml==6.1.3",
            "python-pptx==1.0.2", "Pillow==11.3.0",
        ]
        subprocess.run(command, check=True)
        if str(DEPENDENCIES) not in sys.path:
            sys.path.append(str(DEPENDENCIES))
        if pandoc.is_file():
            return str(pandoc)
    raise RuntimeError("无法找到 Pandoc。")


def run_pandoc(pandoc: str, source: str, source_format: str,
               target_format: str, extra: list[str] | None = None,
               output: Path | None = None) -> str:
    """用 UTF-8 文本经 Pandoc 转换；输出 DOCX 时写入指定文件。"""
    command = [pandoc, "--from", source_format, "--to", target_format]
    command.extend(extra or [])
    if output is not None:
        command.extend(["--output", str(output)])
    result = subprocess.run(
        command, input=source.encode("utf-8"),
        stdout=subprocess.PIPE, stderr=subprocess.PIPE, check=False,
    )
    stderr = result.stderr.decode("utf-8", errors="replace")
    if result.returncode:
        raise RuntimeError(f"Pandoc 转换失败：\n{stderr.strip()}")
    if "Could not convert TeX math" in stderr:
        raise ValueError(f"有公式无法转换为可编辑的 Word 公式：\n{stderr.strip()}")
    if stderr.strip():
        print(stderr.strip(), file=sys.stderr)
    return result.stdout.decode("utf-8")


def _as_json(pandoc: str, source: str, source_format: str,
             extra: list[str] | None = None) -> dict:
    return json.loads(run_pandoc(pandoc, source, source_format, "json", extra))


def _resolve_source(raw: str, base: Path) -> Path:
    candidate = Path(raw.replace("\\", "/"))
    if not candidate.is_absolute():
        candidate = base / candidate
    candidate = candidate.resolve()
    if not candidate.is_file():
        raise FileNotFoundError(f"资源文件不存在：{candidate}")
    return candidate


def _flat_opc_to_docx(source: Path, target: Path) -> None:
    """将 Word 导出的 Flat OPC 副本还原成 Pandoc 可读取的 DOCX。"""
    from lxml import etree

    package_ns = "http://schemas.microsoft.com/office/2006/xmlPackage"
    types_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    package = etree.parse(str(source)).getroot()
    types = etree.Element(f"{{{types_ns}}}Types", nsmap={None: types_ns})
    with ZipFile(target, "w", compression=ZIP_DEFLATED) as archive:
        for part in package:
            name = part.get(f"{{{package_ns}}}name")
            content_type = part.get(f"{{{package_ns}}}contentType")
            if not name or not name.startswith("/") or not content_type or len(part) != 1:
                raise ValueError(f"Word 版式副本包含无效的包部件：{name!r}")
            payload = part[0]
            if payload.tag == f"{{{package_ns}}}xmlData" and len(payload) == 1:
                data = etree.tostring(payload[0], encoding="UTF-8", xml_declaration=True)
            elif payload.tag == f"{{{package_ns}}}binaryData":
                data = base64.b64decode("".join(payload.itertext()))
            else:
                raise ValueError(f"Word 版式副本包含无法识别的包部件：{name}")
            archive.writestr(name.lstrip("/"), data)
            etree.SubElement(types, f"{{{types_ns}}}Override",
                             PartName=name, ContentType=content_type)
        archive.writestr("[Content_Types].xml",
                         etree.tostring(types, encoding="UTF-8", xml_declaration=True))


def _snapshot_operation_reference(reference: Path, directory: Path) -> Path:
    """复制版式文件；若被 Word 独占打开，则读取当前打开文档的副本。"""
    snapshot = directory / "operation_reference.docx"
    try:
        shutil.copyfile(reference, snapshot)
        return snapshot
    except PermissionError as error:
        flat_opc = directory / "operation_reference.xml"
        helper = ROOT / "转换工具" / "snapshot_open_word.ps1"
        result = subprocess.run(
            ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass",
             "-File", str(helper), "-ReferencePath", str(reference),
             "-OutputPath", str(flat_opc)],
            stdout=subprocess.PIPE, stderr=subprocess.PIPE, check=False,
        )
        if result.returncode:
            details = result.stderr.decode("utf-8", errors="replace").strip()
            raise PermissionError(
                f"无法读取正在使用的 Word 版式文件：{reference}。"
                "请在 Word 中关闭它，或用 --reference 指定可读取的副本。"
                + (f"\n{details}" if details else "")
            ) from error
        _flat_opc_to_docx(flat_opc, snapshot)
        print(f"提示：{reference} 正在 Word 中打开，已使用当前打开文档的版式副本。",
              file=sys.stderr)
        return snapshot


def _xlsx_to_html(spec: dict, base: Path) -> str:
    """把一个 XLSX 工作表区域转换成保留合并关系的 HTML 表格。"""
    from openpyxl import load_workbook
    from openpyxl.utils.cell import range_boundaries

    path = _resolve_source(spec["path"], base)
    workbook = load_workbook(path, data_only=False, read_only=False)
    values = load_workbook(path, data_only=True, read_only=False)
    try:
        sheet_name = spec.get("sheet") or workbook.sheetnames[0]
        sheet = workbook[sheet_name]
        value_sheet = values[sheet_name]
        region = spec.get("range") or sheet.calculate_dimension()
        min_col, min_row, max_col, max_row = range_boundaries(region)
        merges = {}
        covered = set()
        for merged in sheet.merged_cells.ranges:
            intersects = not (
                merged.max_col < min_col or merged.min_col > max_col
                or merged.max_row < min_row or merged.min_row > max_row
            )
            if not intersects:
                continue
            inside = (min_col <= merged.min_col <= merged.max_col <= max_col
                      and min_row <= merged.min_row <= merged.max_row <= max_row)
            if not inside:
                raise ValueError(f"XLSX 区域 {region} 截断了合并单元格 {merged}")
            merges[(merged.min_row, merged.min_col)] = (
                merged.max_row - merged.min_row + 1,
                merged.max_col - merged.min_col + 1,
            )
            for row in range(merged.min_row, merged.max_row + 1):
                for col in range(merged.min_col, merged.max_col + 1):
                    if (row, col) != (merged.min_row, merged.min_col):
                        covered.add((row, col))
        parts = ["<table>"]
        if spec.get("caption"):
            parts.append(f"<caption>{html.escape(str(spec['caption']))}</caption>")
        for row in range(min_row, max_row + 1):
            parts.append("<tr>")
            for col in range(min_col, max_col + 1):
                if (row, col) in covered:
                    continue
                tag = "th" if row == min_row and spec.get("header", True) else "td"
                rowspan, colspan = merges.get((row, col), (1, 1))
                attrs = (f' rowspan="{rowspan}"' if rowspan > 1 else "")
                attrs += f' colspan="{colspan}"' if colspan > 1 else ""
                visible = value_sheet.cell(row, col).value
                value = visible if visible is not None else sheet.cell(row, col).value
                content = "" if value is None else html.escape(str(value))
                parts.append(f"<{tag}{attrs}>{content}</{tag}>")
            parts.append("</tr>")
        parts.append("</table>")
        return "".join(parts)
    finally:
        workbook.close()
        values.close()


def _prepare_markdown(source: str, base: Path) -> tuple[str, dict[str, str]]:
    tables = {}

    def replace_table(match: re.Match, markup: str) -> str:
        token = f"MDCONVERTTABLE{len(tables) + 1:06d}"
        tables[token] = markup
        return f"\n\n{token}\n\n"

    def xlsx_match(match: re.Match) -> str:
        try:
            spec = json.loads(match.group(1))
        except json.JSONDecodeError as exc:
            raise ValueError(f"XLSX 表格指令不是有效 JSON：{match.group(0)}") from exc
        return replace_table(match, _xlsx_to_html(spec, base))

    source = XLSX_DIRECTIVE.sub(xlsx_match, source)
    source = HTML_TABLE.sub(lambda match: replace_table(match, match.group(0)), source)

    def image_match(match: re.Match) -> str:
        path = _resolve_source(match.group(2), base).as_posix()
        return match.group(1) + f"<{path}>" + match.group(3)

    source = MARKED_IMAGE.sub(image_match, source)
    source = STANDALONE_IMAGE.sub(
        lambda match: f'![](<{_resolve_source(match.group(1), base).as_posix()}>)',
        source,
    )
    source = re.sub(
        r"(?<![\\$])\$(?!\$)([^$\n]*?)(?<![\\$])\$(?!\$)",
        lambda match: "$" + match.group(1).strip() + "$",
        source,
    )
    def align_multiline_math(match: re.Match) -> str:
        original_expression = match.group(1).strip()
        expression = original_expression
        if r"\begin{" in expression:
            return match.group(0)
        if len(expression) > 180 and r"\qquad" in expression:
            expression = expression.replace(r"\qquad", r"\\", 1)
        lines = []
        depth = 0
        segment_start = 0
        index = 0
        while index < len(expression):
            if expression[index] == "{":
                depth += 1
            elif expression[index] == "}":
                depth -= 1
            elif depth == 0 and expression.startswith(r"\\", index):
                lines.append(expression[segment_start:index].strip())
                index += 2
                segment_start = index
                continue
            index += 1
        if lines:
            lines.append(expression[segment_start:].strip())
            lines = ["& " + line for line in lines]
            return "$$\n\\begin{aligned}\n" + (r" \\" + "\n").join(lines) + "\n\\end{aligned}\n$$"
        return match.group(0)

    source = re.sub(r"\$\$([\s\S]*?)\$\$", align_multiline_math, source)
    return source, tables


def _split_softbreak_paragraphs(document: dict) -> None:
    """把源稿相邻正文行视为独立段落，保留列表与表格结构。"""
    def split_blocks(blocks: list[dict]) -> list[dict]:
        result = []
        for block in blocks:
            kind = block["t"]
            if kind in ("Para", "Plain"):
                pieces = [[]]
                for inline in block["c"]:
                    if inline.get("t") == "SoftBreak":
                        if pieces[-1]:
                            pieces.append([])
                    else:
                        pieces[-1].append(inline)
                result.extend({"t": kind, "c": part} for part in pieces if part)
                continue
            if kind == "Div":
                block["c"][1] = split_blocks(block["c"][1])
            elif kind == "BlockQuote":
                block["c"] = split_blocks(block["c"])
            elif kind == "BulletList":
                block["c"] = [split_blocks(item) for item in block["c"]]
            elif kind == "OrderedList":
                block["c"][1] = [split_blocks(item) for item in block["c"][1]]
            result.append(block)
        return result

    document["blocks"] = split_blocks(document["blocks"])


def _remove_duplicate_captions(document: dict) -> None:
    """消除 Word 回转时紧邻图表重复出现的同名图题或表题。"""
    def inline_text(inlines: list[dict]) -> str:
        pieces = []
        for inline in inlines:
            if not isinstance(inline, dict):
                continue
            kind = inline.get("t")
            if kind == "Str":
                pieces.append(inline["c"])
            elif kind in ("Space", "SoftBreak"):
                pieces.append(" ")
            elif isinstance(inline.get("c"), list):
                pieces.append(inline_text(inline["c"]))
        return "".join(pieces).strip()

    def block_text(block: dict) -> str:
        return inline_text(block["c"]) if block["t"] in ("Para", "Plain") else ""

    blocks = document["blocks"]
    remove = set()
    for index, block in enumerate(blocks):
        if block["t"] not in ("Table", "Figure"):
            continue
        caption_blocks = block["c"][1][1]
        caption = " ".join(block_text(item) for item in caption_blocks).strip()
        if not caption:
            continue
        for neighbor in (index - 1, index + 1):
            if 0 <= neighbor < len(blocks) and block_text(blocks[neighbor]) == caption:
                remove.add(neighbor)
    document["blocks"] = [block for index, block in enumerate(blocks) if index not in remove]


def _insert_tables(document: dict, tables: dict[str, str], pandoc: str) -> None:
    parsed = {}
    for token, markup in tables.items():
        table_doc = _as_json(pandoc, markup, "html")
        candidates = [block for block in table_doc["blocks"] if block["t"] == "Table"]
        if len(candidates) != 1:
            raise ValueError(f"无法把表格转换为 Word 原生表格：{token}")
        parsed[token] = candidates[0]

    def walk_blocks(blocks: list[dict]) -> list[dict]:
        output = []
        for block in blocks:
            if block["t"] in ("Para", "Plain"):
                inlines = block["c"]
                if len(inlines) == 1 and inlines[0].get("t") == "Str":
                    token = inlines[0]["c"]
                    if token in parsed:
                        output.append(copy.deepcopy(parsed[token]))
                        continue
            if block["t"] in ("Div", "BlockQuote"):
                index = 1 if block["t"] == "Div" else None
                if index is None:
                    block["c"] = walk_blocks(block["c"])
                else:
                    block["c"][index] = walk_blocks(block["c"][index])
            output.append(block)
        return output

    document["blocks"] = walk_blocks(document["blocks"])
    if any(token in json.dumps(document, ensure_ascii=False) for token in parsed):
        raise ValueError("有表格占位符没有被识别，请检查其前后是否单独成段。")


def _walk_nodes(value, transform):
    if isinstance(value, list):
        return [_walk_nodes(item, transform) for item in value]
    if isinstance(value, dict):
        replacement = transform(value)
        if replacement is not value:
            return replacement
        if "c" in value:
            value["c"] = _walk_nodes(value["c"], transform)
    return value


def _citation_keys(document: dict) -> list[str]:
    keys = []

    def collect(node: dict):
        if node.get("t") == "Cite":
            for citation in node["c"][0]:
                key = citation["citationId"]
                if key not in keys:
                    keys.append(key)
        return node

    _walk_nodes(document["blocks"], collect)
    return keys


def _encode_references(document: dict, keys: list[str], figure_numbers: dict[str, int]) -> None:
    """将文献与按图名引用的编号写入 Word 字符样式, 回转时恢复原始引用语法."""
    order = {key: index + 1 for index, key in enumerate(keys)}

    def convert(node: dict):
        if node.get("t") == "Str" and "{{figref:" in node["c"]:
            text = node["c"]
            inlines = []
            cursor = 0
            for match in FIGURE_REFERENCE.finditer(text):
                name, panels = match.groups()
                names = name.split(",")
                for figure_name in names:
                    if figure_name not in figure_numbers:
                        raise ValueError(f"正文引用的图名没有插图定义：{figure_name}")
                if len(names) > 1 and panels:
                    raise ValueError("多个图片的组合引用不能附加面板；请分别引用各图的面板")
                if match.start() > cursor:
                    inlines.append({"t": "Str", "c": text[cursor:match.start()]})
                literal = match.group(0)
                marker = base64.urlsafe_b64encode(zlib.compress(literal.encode("utf-8"))).decode("ascii").rstrip("=")
                if len(names) == 1:
                    visible = f"Fig. {figure_numbers[name]}"
                else:
                    ranges = []
                    start = end = figure_numbers[names[0]]
                    for figure_name in names[1:]:
                        number = figure_numbers[figure_name]
                        if number == end + 1:
                            end = number
                        else:
                            ranges.append(str(start) if start == end else f"{start}–{end}")
                            start = end = number
                    ranges.append(str(start) if start == end else f"{start}–{end}")
                    visible = "Figs. " + ", ".join(ranges)
                if panels:
                    visible += panels.replace("-", "–").replace(",", ", ")
                inlines.append({"t": "Span", "c": [
                    ["", [], [["custom-style", FIGURE_REFERENCE_STYLE_PREFIX + marker]]],
                    [{"t": "Str", "c": visible}],
                ]})
                cursor = match.end()
            inlines.append({"t": "Str", "c": text[cursor:]})
            if any("{{figref:" in item.get("c", "") for item in inlines if item["t"] == "Str"):
                raise ValueError("图片正文引用格式无效；使用 {{figref:图名}} 或 {{figref:图名|a,b}}")
            return {"t": "Span", "c": [["", [], []], inlines]}
        if node.get("t") != "Cite":
            return node
        cited = [entry["citationId"] for entry in node["c"][0]]
        literal = "[@" + "; @".join(cited) + "]"
        marker = base64.urlsafe_b64encode(zlib.compress(literal.encode("utf-8"))).decode("ascii").rstrip("=")
        style = CITATION_STYLE_PREFIX + marker
        visible = ",".join(str(order[key]) for key in cited)
        return {
            "t": "Span", "c": [
                ["", [], [["custom-style", style]]],
                [{"t": "Superscript", "c": [{"t": "Str", "c": visible}]}],
            ],
        }

    document["blocks"] = _walk_nodes(document["blocks"], convert)


def _append_bibliography(document: dict, original: dict, pandoc: str,
                         bibliography: Path, keys: list[str]) -> None:
    if not keys:
        return
    if not bibliography.is_file():
        raise FileNotFoundError(f"引用了 .bib 键，但找不到文献库：{bibliography}")
    processed = _as_json(
        pandoc, json.dumps(original, ensure_ascii=False), "json",
        ["--citeproc", f"--bibliography={bibliography}", f"--csl={CSL}"],
    )
    reference_divs = [
        block for block in processed["blocks"]
        if block.get("t") == "Div" and block["c"][0][0] == "refs"
    ]
    if len(reference_divs) != 1:
        raise RuntimeError("无法生成参考文献列表，请检查 .bib 中的引用键。")
    reference_entries = reference_divs[0]["c"][1]
    if len(reference_entries) != len(keys):
        raise ValueError("部分引用键没有在 .bib 中找到，参考文献数量与引用键不一致。")
    heading = {"t": "Para", "c": [{"t": "Strong", "c": [{"t": "Str", "c": "References"}]}]}
    paragraphs = [heading]
    for entry in reference_entries:
        paragraphs.extend(entry["c"][1])
    for paragraph in paragraphs:
        if paragraph["t"] not in ("Para", "Plain"):
            continue
        document["blocks"].append({
            "t": "Div", "c": [
                ["", [], [["custom-style", BIBLIOGRAPHY_STYLE]]],
                [paragraph],
            ],
        })


def _format_content_paragraphs(document, body_spacing: float, body_size: float) -> None:
    """统一正文段落节奏，并使块级公式保持独立、居中。"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt

    for paragraph in document.paragraphs:
        paragraph_format = paragraph.paragraph_format
        if paragraph._p.find(qn("m:oMathPara")) is not None:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph_format.first_line_indent = Pt(0)
            paragraph_format.line_spacing = 1.0
            paragraph_format.space_before = Pt(9)
            paragraph_format.space_after = Pt(11)
            paragraph_format.keep_together = True
        elif paragraph.style.name.startswith("Heading"):
            paragraph_format.first_line_indent = Pt(0)
            paragraph_format.keep_with_next = True
        elif paragraph.style.name in ("Body Text", "First Paragraph", "Normal"):
            if paragraph.text.strip():
                paragraph_format.first_line_indent = Pt(2 * body_size)
                paragraph_format.line_spacing = body_spacing
                paragraph_format.space_after = Pt(6)
        elif paragraph.style.name in ("Figure Caption", "Image Caption"):
            paragraph_format.first_line_indent = Pt(0)
            paragraph_format.keep_together = True
        elif paragraph.style.name == "Table Caption":
            paragraph_format.first_line_indent = Pt(0)
            paragraph_format.keep_with_next = True
        if paragraph._p.find(".//" + qn("w:drawing")) is not None:
            paragraph_format.first_line_indent = Pt(0)
            paragraph_format.keep_with_next = True


def _layout_images(document, alignment: str, sizing: str) -> None:
    """按版心等比放置独立图片，并保留行内图片的原有排版。"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    alignments = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }
    section = document.sections[0]
    max_width = section.page_width - section.left_margin - section.right_margin
    max_height = int((section.page_height - section.top_margin - section.bottom_margin) * 0.85)
    paragraphs = {id(item._p): item for item in document.paragraphs}
    for shape in document.inline_shapes:
        paragraph_element = shape._inline.getparent()
        while paragraph_element is not None and paragraph_element.tag != qn("w:p"):
            paragraph_element = paragraph_element.getparent()
        if paragraph_element is None:
            continue
        paragraph = paragraphs.get(id(paragraph_element))
        if (paragraph is None or paragraph.text.strip()
                or len(paragraph._p.findall(".//" + qn("wp:inline"))) != 1):
            continue
        paragraph.alignment = alignments[alignment]
        if (sizing == "fit" and shape.width and shape.height
                and not shape._inline.docPr.get("title", "").startswith(PPT_EDITED_TITLE_PREFIX)):
            width, height = shape.width, shape.height
            if max_width * height <= max_height * width:
                shape.width = max_width
                shape.height = round(height * max_width / width)
            else:
                shape.width = round(width * max_height / height)
                shape.height = max_height


def _remove_heading_numbering(document) -> None:
    """清除参考模板附着在标题样式上的自动章节编号。"""
    from docx.oxml.ns import qn

    heading_ids = set()
    for level in range(1, 7):
        name = f"Heading {level}"
        if name not in document.styles:
            continue
        style = document.styles[name]
        heading_ids.add(style.style_id)
        if style.element.pPr is not None and style.element.pPr.numPr is not None:
            style.element.pPr.remove(style.element.pPr.numPr)
    for paragraph in document.paragraphs:
        if paragraph.style.style_id not in heading_ids:
            continue
        if paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None:
            paragraph._p.pPr.remove(paragraph._p.pPr.numPr)
    numbering = document.part.numbering_part.element
    for level in numbering.iter(qn("w:lvl")):
        linked_style = level.find(qn("w:pStyle"))
        if linked_style is not None and linked_style.get(qn("w:val")) in heading_ids:
            level.remove(linked_style)


def _fit_long_equations(document, size_half_points: int) -> None:
    """仅缩小超长块级公式，避免公式内容伸入页边距。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for equation in document.element.body.iter(qn("m:oMathPara")):
        text_length = sum(len(item.text or "") for item in equation.iter(qn("m:t")))
        if text_length < 120:
            continue
        for math_run in equation.iter(qn("m:r")):
            properties = math_run.find(qn("w:rPr"))
            if properties is None:
                properties = OxmlElement("w:rPr")
                math_run.insert(0, properties)
            size = properties.find(qn("w:sz"))
            if size is None:
                size = OxmlElement("w:sz")
                properties.append(size)
            size.set(qn("w:val"), str(size_half_points))


def _repair_underbraces(document) -> None:
    """把 Pandoc 写出的固定宽度 ⏟ 转成 Word 可伸展的下括号。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for inner in list(document.element.body.iter(qn("m:limLow"))):
        base = inner.find(qn("m:e"))
        limit = inner.find(qn("m:lim"))
        if base is None or limit is None:
            continue
        marker = "".join(item.text or "" for item in limit.iter(qn("m:t")))
        if marker != "⏟":
            continue
        group = OxmlElement("m:groupChr")
        properties = OxmlElement("m:groupChrPr")
        character = OxmlElement("m:chr")
        character.set(qn("m:val"), "⏟")
        position = OxmlElement("m:pos")
        position.set(qn("m:val"), "bot")
        properties.extend((character, position))
        group.append(properties)
        inner.remove(base)
        group.append(base)
        inner.getparent().replace(inner, group)


def _apply_nature_layout(path: Path, image_align: str, image_size: str) -> None:
    """设置 Nature 初次投稿需要的字体、双倍行距、行号和表格分页。"""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    document = Document(path)
    for section in document.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        sect_pr = section._sectPr
        for old in sect_pr.findall(qn("w:lnNumType")):
            sect_pr.remove(old)
        numbering = OxmlElement("w:lnNumType")
        numbering.set(qn("w:countBy"), "1")
        numbering.set(qn("w:start"), "1")
        numbering.set(qn("w:restart"), "continuous")
        sect_pr.append(numbering)
    for style in document.styles:
        if style.type != 1:
            continue
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.line_spacing = 2
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(6)
    if "Image Caption" in document.styles:
        document.styles["Image Caption"].font.italic = False
    for level, size in ((1, 14), (2, 12), (3, 12), (4, 12), (5, 12), (6, 12)):
        name = f"Heading {level}"
        if name not in document.styles:
            continue
        style = document.styles[name]
        style.font.size = Pt(size)
        style.font.bold = level <= 3
        style.font.italic = level >= 4
        style.paragraph_format.space_before = Pt(12 if level <= 2 else 9)
        style.paragraph_format.space_after = Pt(6)
    for paragraph in document.paragraphs:
        if paragraph._p.find(qn("m:oMathPara")) is None:
            paragraph.paragraph_format.line_spacing = 2
        if paragraph.style.name == "Table Caption":
            paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            if run.font.name is None:
                run.font.name = "Times New Roman"
            if run.font.size is None:
                run.font.size = Pt(12)
            if paragraph.style.name == "Table Caption":
                run.font.bold = True
                run.font.italic = False
    _format_content_paragraphs(document, 2.0, 12)
    _layout_images(document, image_align, image_size)
    _fit_long_equations(document, 20)
    _repair_underbraces(document)
    body = document.element.body
    elements = list(body)
    if elements and elements[0].tag == qn("w:p"):
        first = elements[0]
        has_content = any(
            next(first.iter(qn(tag)), None) is not None
            for tag in ("w:t", "w:drawing", "m:oMath")
        )
        if not has_content:
            body.remove(first)
            elements.pop(0)

    for index in range(len(elements) - 2, -1, -1):
        if elements[index].tag != qn("w:tbl"):
            continue
        following = elements[index + 1]
        if following.tag != qn("w:p"):
            continue
        has_content = any(
            next(following.iter(qn(tag)), None) is not None
            for tag in ("w:t", "w:drawing", "m:oMath", "w:br")
        )
        if not has_content:
            body.remove(following)
            elements.pop(index + 1)

    def page_break_before(paragraph_element) -> None:
        properties = paragraph_element.find(qn("w:pPr"))
        if properties is None:
            properties = OxmlElement("w:pPr")
            paragraph_element.insert(0, properties)
        if properties.find(qn("w:pageBreakBefore")) is None:
            properties.append(OxmlElement("w:pageBreakBefore"))

    for index, element in enumerate(elements):
        if element.tag != qn("w:tbl"):
            continue
        previous = elements[index - 1] if index else None
        caption = None
        if previous is not None and previous.tag == qn("w:p"):
            style = previous.find(".//" + qn("w:pStyle"))
            if style is not None and style.get(qn("w:val")) == "TableCaption":
                caption = previous
        if caption is not None:
            page_break_before(caption)
        else:
            first_paragraph = element.find(".//" + qn("w:p"))
            if first_paragraph is not None:
                page_break_before(first_paragraph)
        for following in elements[index + 1:]:
            if following.tag == qn("w:p"):
                has_text = bool(following.findall(".//" + qn("w:t")))
                has_drawing = bool(following.findall(".//" + qn("w:drawing")))
                if not (has_text or has_drawing):
                    continue
                page_break_before(following)
                break
            if following.tag == qn("w:tbl"):
                break
    document.save(path)


def _apply_operation_layout(path: Path, reference: Path, image_align: str, image_size: str) -> None:
    """沿用 Operation 页框和标题层级，形成适合通读的 Word 稿。"""
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    source = Document(reference)
    document = Document(path)
    # 本机 Word 模板可能只保留本地化样式编号; 补齐 Pandoc 写出的图注和字符样式依赖, 保留回转语义.
    for name in ("Body Text Char", "Verbatim Char"):
        if name not in document.styles:
            style = document.styles.add_style(name, WD_STYLE_TYPE.CHARACTER)
            style.base_style = document.styles["Default Paragraph Font"]
            if name == "Verbatim Char":
                style.font.name = "Courier New"
    if "Image Caption" not in document.styles:
        style = document.styles.add_style("Image Caption", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles["Caption"]
    document.styles["Image Caption"].font.italic = False
    example = source.sections[0]
    for section in document.sections:
        for property_name in (
            "page_width", "page_height", "top_margin", "bottom_margin",
            "left_margin", "right_margin", "header_distance", "footer_distance",
        ):
            setattr(section, property_name, getattr(example, property_name))
    for name in ("Normal", "Body Text", "First Paragraph"):
        if name not in document.styles:
            document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style = document.styles[name]
        style.font.size = Pt(10.5)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.first_line_indent = Pt(21)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    heading_rules = (
        (1, 22, 17, 16.5), (2, 16, 13, 13), (3, 14, 13, 10),
        (4, 12, 11, 8), (5, 11, 10, 7), (6, 10.5, 9, 6),
    )
    for level, size, before, after in heading_rules:
        name = f"Heading {level}"
        if name not in document.styles:
            document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style = document.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.line_spacing = 1.25
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
        )
    _format_content_paragraphs(document, 1.5, 10.5)
    paragraphs = document.paragraphs
    for index, paragraph in enumerate(paragraphs):
        properties = paragraph._p.pPr
        is_table_caption = (properties is not None and properties.pStyle is not None
                            and properties.pStyle.val == "TableCaption")
        # 原生公式不计入 paragraph.text; 优先用图注样式识别, 避免公式图注回退为正文格式.
        is_figure_caption = paragraph.style.name in ("Figure Caption", "Image Caption")
        if not is_figure_caption and index and paragraphs[index - 1].text.strip() == "":
            descriptions = [item.get("descr", "").strip()
                            for item in paragraphs[index - 1]._p.xpath(".//wp:docPr")]
            is_figure_caption = bool(paragraph.text.strip() and paragraph.text.strip() in descriptions)
        if not (is_table_caption or is_figure_caption):
            continue
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.keep_together = True
        if is_table_caption:
            paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            run.font.name = "华文仿宋"
            run.font.size = Pt(9)
            run._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "华文仿宋")
    _remove_heading_numbering(document)
    _layout_images(document, image_align, image_size)
    _fit_long_equations(document, 19)
    _repair_underbraces(document)
    document.save(path)


def markdown_to_word(source: Path, target: Path, bibliography: Path | None = None,
                     profile: str = "nature", reference: Path | None = None,
                     image_align: str = "center", image_size: str = "fit",
                     english: bool = False) -> None:
    """在 temp 中生成 Word 后发布; english 选择 PPT 双语备注的英文部分, 无分隔标记时保留原图注."""
    pandoc = ensure_dependencies()
    from docx import Document
    from PIL import Image

    if image_align not in ("left", "center", "right"):
        raise ValueError(f"未知图片对齐方式：{image_align}")
    if image_size not in ("fit", "original"):
        raise ValueError(f"未知图片放置方式：{image_size}")

    source = source.resolve()
    target = target.resolve()
    if source == target:
        raise ValueError("输入和输出不能是同一个文件。")
    TEMP_ROOT.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="manuscript_figures_", dir=TEMP_ROOT) as temporary:
        figure_root = Path(temporary)
        layout_reference = None
        if profile == "operation":
            reference = (reference or OPERATION_REFERENCE).resolve()
            if not reference.is_file():
                raise FileNotFoundError(f"找不到 Operation 版式文件：{reference}")
            layout_reference = _snapshot_operation_reference(reference, figure_root)
        elif profile != "nature":
            raise ValueError(f"未知 Word 版式：{profile}")
        figure_specs = {}
        figure_index = {}
        figure_owners = {}
        deck_snapshots = {}
        rendered = {}
        export_requests = []
        source_text = source.read_text(encoding="utf-8")
        # 图名到 1 起始编号的映射, 按插图位置排序; Word 内改图后的来源注释也保留编号资格.
        definitions = [(match.start(), match.group(2).strip()) for match in PPT_FIGURE.finditer(source_text)]
        definitions.extend((match.start(), json.loads(match.group(1))["name"])
                           for match in PPT_EDITED_SOURCE.finditer(source_text))
        figure_numbers = {}
        for _, name in sorted(definitions):
            if name in figure_numbers:
                raise ValueError(f"同一图名不能重复插图：{name}；再次提及时使用 {{{{figref:{name}}}}}")
            figure_numbers[name] = len(figure_numbers) + 1
        preflight = _as_json(pandoc, source_text, PANDOC_INPUT)
        _encode_references(preflight, _citation_keys(preflight), figure_numbers)
        pieces = []
        cursor = 0
        for match in PPT_FIGURE.finditer(source_text):
            ppt_label, figure_name = match.group(1), match.group(2).strip()
            if not re.fullmatch(r"[^\s|{},]+", figure_name):
                raise ValueError(f"PPT 图名无效：{figure_name!r}")
            ppt_path = _resolve_source(ppt_label, source.parent)
            if ppt_path not in figure_index:
                snapshot = figure_root / f"deck_{len(deck_snapshots) + 1:06d}.pptx"
                shutil.copyfile(ppt_path, snapshot)
                deck_snapshots[ppt_path] = snapshot
                figure_index[ppt_path] = read_ppt_figures(snapshot)
                for name in figure_index[ppt_path]:
                    if name in figure_owners and figure_owners[name] != ppt_path:
                        raise ValueError(f"跨 PPT 图名重复：{name}，来自 {figure_owners[name]} 与 {ppt_path}")
                    figure_owners[name] = ppt_path
            if figure_name not in figure_index[ppt_path]:
                raise ValueError(f"PPT 图名不存在：{ppt_path} 中的 {figure_name}")
            page, caption = figure_index[ppt_path][figure_name]
            # 独立的 ENGLISH: 或 ENGLISH：行分隔中文与英文图注; 匹配实际换行, 不匹配字面的 \\n.
            language_marker = re.search(r"(?:\A|(?<=[\r\n]))[ \t]*ENGLISH[:：][ \t]*(?:\r?\n|\r)", caption)
            if language_marker is not None:
                caption = caption[language_marker.end():] if english else caption[:language_marker.start()]
            caption = caption.strip()
            caption = MANUAL_FIGURE_NUMBER.sub("", caption)
            if caption:
                title_end = re.search(r"。|\.(?=\s|$)|\n", caption)
                split = title_end.end() if title_end else len(caption)
                # 只加粗图注段首的面板字母或范围, 如 a-c、a–c; 保留原分隔符、逗号、正文、公式及已有 Markdown 加粗标记.
                body = re.sub(r"(?m)^([ \t]*)([a-z](?:[ \t]*[-–][ \t]*[a-z])?)(?=[,，、])", r"\1**\2**", caption[split:])
                caption = f"**Fig. {figure_numbers[figure_name]} | {caption[:split]}**{body}"
            if not caption:
                print(f"提示：{ppt_path} 第 {page} 页备注为空，Word 中不会生成该图的图注。", file=sys.stderr)
            figure_key = (ppt_path, figure_name)
            if figure_key not in rendered:
                image = figure_root / f"figure_{len(rendered) + 1:06d}.png"
                rendered[figure_key] = image
                export_requests.append((deck_snapshots[ppt_path], page, image))
            image = rendered[figure_key]
            token = f"MDPPTFIGTOKEN{len(figure_specs) + 1:06d}"
            figure_specs[token] = {
                "path": ppt_label, "resolved": str(ppt_path), "name": figure_name,
                "caption": caption, "image": image.as_posix(),
            }
            pieces.extend((source_text[cursor:match.start()], f"\n\n{token}\n\n"))
            cursor = match.end()
        export_ppt_figures(export_requests, TEMP_ROOT)
        image_widths = {}
        for image in rendered.values():
            with Image.open(image) as bitmap:
                image_widths[image.as_posix()] = min(5.8, 7.4 * bitmap.width / bitmap.height)
        for spec in figure_specs.values():
            spec["width"] = image_widths[spec["image"]]
        pieces.append(source_text[cursor:])
        expanded_source = "".join(pieces)
        if re.search(r"(?m)^[ \t]*\{\{pptfig:", expanded_source):
            raise ValueError('PPT 图集引用格式无效；应独立成行写 {{pptfig:"PPT路径"|图名}}')
        text, tables = _prepare_markdown(expanded_source, source.parent)
        document = _as_json(pandoc, text, PANDOC_INPUT, [f"--resource-path={source.parent}"])
        _insert_tables(document, tables, pandoc)
        replacements = []
        recognized_figures = set()
        edited_origin = None
        for block in document["blocks"]:
            if block.get("t") == "RawBlock" and block["c"][0] == "html":
                edited_match = PPT_EDITED_COMMENT.fullmatch(block["c"][1].strip())
                if edited_match:
                    edited_origin = json.loads(edited_match.group(1))
                    continue
            if edited_origin is not None:
                image_nodes = []
                pending = [block]
                while pending:
                    item = pending.pop()
                    if isinstance(item, dict):
                        if item.get("t") == "Image":
                            image_nodes.append(item)
                        elif "c" in item:
                            pending.append(item["c"])
                    elif isinstance(item, list):
                        pending.extend(reversed(item))
                if len(image_nodes) != 1:
                    raise ValueError("pptfig-edited 来源注释后必须紧跟一张图片")
                origin_path = Path(edited_origin["path"].replace("\\", "/"))
                resolved_origin = origin_path if origin_path.is_absolute() else source.parent / origin_path
                metadata = {"path": edited_origin["path"], "resolved": str(resolved_origin.resolve()), "name": edited_origin["name"]}
                if block.get("t") == "Figure" and block["c"][1][1]:
                    caption_document = {"pandoc-api-version": document["pandoc-api-version"], "meta": {}, "blocks": block["c"][1][1]}
                    caption_text = run_pandoc(pandoc, json.dumps(caption_document, ensure_ascii=False), "json", PANDOC_OUTPUT, ["--wrap=none"])
                    caption_text = re.sub(r"(?m)^(\*{0,2}Fig\.\s*)\d+(\s*\\?\|)",
                                          rf"\g<1>{figure_numbers[edited_origin['name']]}\g<2>", caption_text)
                    block["c"][1][1] = _as_json(pandoc, caption_text, PANDOC_INPUT)["blocks"]
                    image_nodes[0]["c"][1] = block["c"][1][1][0]["c"]
                marker = base64.urlsafe_b64encode(zlib.compress(json.dumps(metadata, ensure_ascii=False).encode("utf-8"))).decode("ascii").rstrip("=")
                image_nodes[0]["c"][2][1] = PPT_EDITED_TITLE_PREFIX + marker
                edited_origin = None
            inlines = block.get("c") if block.get("t") in ("Para", "Plain") else None
            token = inlines[0].get("c") if isinstance(inlines, list) and len(inlines) == 1 and inlines[0].get("t") == "Str" else None
            if token not in figure_specs:
                replacements.append(block)
                continue
            spec = figure_specs[token]
            recognized_figures.add(token)
            caption_blocks = _as_json(pandoc, spec["caption"], PANDOC_INPUT)["blocks"] if spec["caption"] else []
            alt = caption_blocks[0]["c"] if caption_blocks and caption_blocks[0]["t"] in ("Para", "Plain") else []
            replacements.append({
                "t": "Figure", "c": [
                    ["", [], []], [None, caption_blocks],
                    [{"t": "Plain", "c": [{"t": "Image", "c": [
                        ["", [], [["width", f"{spec['width']:.3f}in"]]], alt,
                        [spec["image"], token],
                    ]}]}],
                ],
            })
        if edited_origin is not None:
            raise ValueError("pptfig-edited 来源注释后缺少图片")
        if recognized_figures != set(figure_specs):
            raise ValueError("PPT 图集引用必须是正文中独立成行的指令，不能写在代码块或其他块内")
        document["blocks"] = replacements
        _split_softbreak_paragraphs(document)
        _remove_duplicate_captions(document)
        original = copy.deepcopy(document)
        keys = _citation_keys(document)
        _encode_references(document, keys, figure_numbers)
        bib = bibliography.resolve() if bibliography else source.parent / "references.bib"
        _append_bibliography(document, original, pandoc, bib, keys)
        target.parent.mkdir(parents=True, exist_ok=True)
        completed_word = figure_root / "completed.docx"
        options = ["--standalone", f"--resource-path={source.parent}"]
        if layout_reference is not None:
            options.append(f"--reference-doc={layout_reference}")
        run_pandoc(pandoc, json.dumps(document, ensure_ascii=False), "json", "docx", options, completed_word)
        if profile == "operation":
            _apply_operation_layout(completed_word, layout_reference, image_align, image_size)
        else:
            _apply_nature_layout(completed_word, image_align, image_size)
        if figure_specs:
            word = Document(completed_word)
            marked = set()
            for shape in word.inline_shapes:
                properties = shape._inline.docPr
                token = properties.get("title")
                if token not in figure_specs:
                    continue
                metadata = {key: figure_specs[token][key] for key in ("path", "resolved", "name")}
                metadata["markup"] = picture_markup(shape._inline)
                encoded = base64.urlsafe_b64encode(zlib.compress(json.dumps(metadata, ensure_ascii=False).encode("utf-8"))).decode("ascii").rstrip("=")
                properties.set("title", PPT_TITLE_PREFIX + encoded)
                marked.add(token)
            if marked != set(figure_specs):
                raise RuntimeError(f"Word 图片标记不完整：{sorted(set(figure_specs) - marked)}")
            word.save(completed_word)
        with ZipFile(completed_word) as package:
            damaged_member = package.testzip()
        if damaged_member is not None:
            raise RuntimeError(f"生成的 Word 压缩包损坏：{damaged_member}")
        try:
            for attempt in range(3):
                try:
                    os.replace(completed_word, target)
                    break
                except OSError as error:
                    if error.errno == errno.EXDEV:
                        shutil.copy2(completed_word, target)
                        break
                    if attempt == 2 or error.errno not in (errno.EACCES, errno.EPERM, errno.EINVAL):
                        raise
                    time.sleep(0.5)
        except OSError as error:
            pending_word = TEMP_ROOT / f"pending_{figure_root.name}.docx"
            shutil.move(completed_word, pending_word)
            raise RuntimeError(
                f"无法写入目标 Word：{target}\n完整 Word 已保留在：{pending_word}\n"
                "请关闭占用目标文件的 Word 窗口，或用 --output 指定另一个文件名。"
            ) from error


def _remove_generated_styles(document: dict, citation_markers: dict[str, str]) -> None:
    """去掉 Word 样式噪音, 保留本工具编码的文献与图名引用."""
    counter = [0]

    def clean(value):
        if isinstance(value, list):
            output = []
            for item in value:
                cleaned = clean(item)
                if (isinstance(item, dict) and item.get("t") in ("Div", "Span")
                        and isinstance(cleaned, list)):
                    output.extend(cleaned)
                elif cleaned is not None or not isinstance(item, dict):
                    output.append(cleaned)
            return output
        if not isinstance(value, dict):
            return value
        kind = value.get("t")
        if kind == "Div":
            attrs, blocks = value["c"]
            style = dict(attrs[2]).get("custom-style", "")
            if style == BIBLIOGRAPHY_STYLE:
                return None
            return clean(blocks)
        if kind == "Span":
            attrs, inlines = value["c"]
            style = dict(attrs[2]).get("custom-style", "")
            prefix = next((prefix for prefix in (CITATION_STYLE_PREFIX, FIGURE_REFERENCE_STYLE_PREFIX)
                           if style.startswith(prefix)), None)
            if prefix is not None:
                encoded = style[len(prefix):]
                padding = "=" * (-len(encoded) % 4)
                literal = zlib.decompress(base64.urlsafe_b64decode(encoded + padding)).decode("utf-8")
                counter[0] += 1
                token = f"MDCITATIONTOKEN{counter[0]:06d}"
                citation_markers[token] = literal
                return {"t": "Str", "c": token}
            return clean(inlines)
        if "c" in value:
            value["c"] = clean(value["c"])
        if kind in ("Table", "Figure"):
            value["c"][0][2] = [entry for entry in value["c"][0][2] if entry[0] != "custom-style"]
        return value

    document["blocks"] = clean(document["blocks"])


def _export_images(document: dict, media_root: Path, destination: Path) -> dict[str, str]:
    image_markers = {}
    images_dir = destination.parent / "manuscript_assets" / "images"

    def export(node: dict):
        if node.get("t") != "Image":
            return node
        source_name = node["c"][2][0]
        candidate = Path(source_name)
        if not candidate.is_file():
            candidate = media_root / source_name
        if not candidate.is_file():
            candidate = media_root / Path(source_name).name
        if not candidate.is_file():
            raise FileNotFoundError(f"无法从 Word 提取图片：{source_name}")
        payload = candidate.read_bytes()
        suffix = candidate.suffix.lower() or ".png"
        images_dir.mkdir(parents=True, exist_ok=True)
        exported = next(
            (item for item in sorted(images_dir.glob(f"image_*{suffix}"))
             if item.read_bytes() == payload),
            None,
        )
        if exported is None:
            number = 1
            while (images_dir / f"image_{number:06d}{suffix}").exists():
                number += 1
            exported = images_dir / f"image_{number:06d}{suffix}"
            shutil.copyfile(candidate, exported)
        token = f"MDIMAGETOKEN{len(image_markers) + 1:06d}"
        image_markers[token] = f'&&"{exported.resolve().as_posix()}"'
        node["c"][2][0] = token
        return node

    document["blocks"] = _walk_nodes(document["blocks"], export)
    return image_markers


def _normalize_equation_spacing(document: dict) -> None:
    """消除 Word 公式回读时在隐形定界符旁累积的空白命令。"""
    def normalize(node: dict):
        if node.get("t") != "Math":
            return node
        expression = node["c"][1]
        expression = re.sub(
            r"(\\(?:left|right)\.)\s+(?:\\\s+)+(?=\\parallel)",
            r"\1 ", expression,
        )
        expression = re.sub(
            r"(\\(?:left|right)\.)(?:\s*\\\s*)*\\parallel\s*",
            r"\1\\parallel", expression,
        )
        expression = re.sub(r"(\\parallel)\s+(?=\\mathbf)", r"\1", expression)
        expression = re.sub(r"(\\right\.)(?:\s*\\\s*)+$", r"\1", expression)
        node["c"][1] = expression
        return node

    document["blocks"] = _walk_nodes(document["blocks"], normalize)


def _normalize_group_char_underbraces(document: dict) -> None:
    """将 Word 伸展下括号回读为稳定的 LaTeX underbrace 写法。"""
    def braced(source: str, position: int) -> tuple[str, int] | None:
        if position >= len(source) or source[position] != "{":
            return None
        depth = 0
        for index in range(position, len(source)):
            if source[index] == "{":
                depth += 1
            elif source[index] == "}":
                depth -= 1
                if depth == 0:
                    return source[position + 1:index], index + 1
        return None

    def replacement(source: str, position: int) -> tuple[str, int] | None:
        label = braced(source, position + len(r"\underset"))
        if label is None:
            return None
        body = braced(source, label[1])
        if body is None or not body[0].startswith(r"\overset"):
            return None
        upper = braced(body[0], len(r"\overset"))
        if upper is None:
            return None
        base = braced(body[0], upper[1])
        if base is None:
            return None
        value = base[0]
        while value.startswith(r"\overset{}"):
            nested = braced(value, len(r"\overset{}"))
            if nested is None:
                return None
            value = nested[0]
        if not value.startswith(r"\underbrace"):
            return None
        brace = braced(value, len(r"\underbrace"))
        if brace is None or brace[0] or brace[1] != len(value):
            return None
        return r"\underbrace{" + upper[0] + "}_{" + label[0] + "}", body[1]

    def normalize(node: dict):
        if node.get("t") != "Math":
            return node
        source = node["c"][1]
        output = []
        cursor = 0
        while True:
            position = source.find(r"\underset", cursor)
            if position < 0:
                output.append(source[cursor:])
                break
            output.append(source[cursor:position])
            parsed = replacement(source, position)
            if parsed is None:
                output.append(r"\underset")
                cursor = position + len(r"\underset")
            else:
                output.append(parsed[0])
                cursor = parsed[1]
        node["c"][1] = "".join(output)
        return node

    document["blocks"] = _walk_nodes(document["blocks"], normalize)


def _normalize_markdown_spacing(markdown: str) -> str:
    """让回读 Markdown 的段落留一个空行，管道表格各行保持相邻。"""
    lines = markdown.splitlines()
    output = []
    in_fence = False
    for index, line in enumerate(lines):
        if re.match(r"^\s*(```|~~~)", line):
            in_fence = not in_fence
        if line.strip() or in_fence:
            output.append(line)
            continue
        previous = next((item for item in reversed(output) if item.strip()), "")
        following = next((item for item in lines[index + 1:] if item.strip()), "")
        if (previous.strip().startswith("|") and previous.strip().endswith("|")
                and following.strip().startswith("|") and following.strip().endswith("|")):
            continue
        if output and output[-1].strip():
            output.append("")
    return "\n".join(output).rstrip() + "\n"


def word_to_markdown(source: Path, target: Path) -> None:
    """把 Word 的文字、公式和表格写回 Markdown, 同时识别未改动的 PPT 图集图片。"""
    pandoc = ensure_dependencies()
    from docx import Document
    from docx.oxml.ns import qn
    source = source.resolve()
    target = target.resolve()
    if source == target:
        raise ValueError("输入和输出不能是同一个文件。")
    TEMP_ROOT.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="manuscript_media_", dir=TEMP_ROOT) as temporary:
        media_root = Path(temporary)
        document = _as_json(
            pandoc, "", "docx+styles",
            [str(source), f"--extract-media={media_root}"],
        )
        citations = {}
        _remove_generated_styles(document, citations)
        _normalize_equation_spacing(document)
        _normalize_group_char_underbraces(document)
        source_word = Document(source)
        for anchor in source_word.element.body.iter(qn("wp:anchor")):
            properties = anchor.find(qn("wp:docPr"))
            title = properties.get("title", "") if properties is not None else ""
            if title.startswith((PPT_TITLE_PREFIX, PPT_EDITED_TITLE_PREFIX)):
                raise ValueError("PPT 图集图片已改为 Word 浮动对象，无法可靠回转；请改为与文字在一行后重试")
        marked_images = {}
        edited_images = {}
        for shape in source_word.inline_shapes:
            inline = shape._inline
            title = inline.docPr.get("title", "")
            if title.startswith(PPT_TITLE_PREFIX):
                prefix = PPT_TITLE_PREFIX
            elif title.startswith(PPT_EDITED_TITLE_PREFIX):
                prefix = PPT_EDITED_TITLE_PREFIX
            else:
                continue
            encoded = title[len(prefix):]
            try:
                metadata = json.loads(zlib.decompress(base64.urlsafe_b64decode(encoded + "=" * (-len(encoded) % 4))))
                relation = inline.xpath(".//a:blip")[0].get(qn("r:embed"))
            except (ValueError, IndexError, KeyError, zlib.error) as exc:
                raise ValueError("Word 中的 PPT 图名标记已损坏") from exc
            if prefix == PPT_TITLE_PREFIX:
                marked_images.setdefault(relation, []).append((metadata, picture_markup(inline), source_word.part.related_parts[relation].blob))
            else:
                edited_images.setdefault(relation, []).append(metadata)
        figure_indices = {}
        deck_snapshots = {}
        rendered_figures = {}
        recovered = {}
        blocks = []
        for block in document["blocks"]:
            image_nodes = []
            pending = [block]
            while pending:
                item = pending.pop()
                if isinstance(item, dict):
                    if item.get("t") == "Image":
                        image_nodes.append(item)
                    elif "c" in item:
                        pending.append(item["c"])
                elif isinstance(item, list):
                    pending.extend(reversed(item))
            if len(image_nodes) != 1:
                blocks.append(block)
                continue
            image_node = image_nodes[0]
            relation = Path(image_node["c"][2][0]).stem
            edited_candidates = edited_images.get(relation, [])
            if edited_candidates:
                edited_origin = edited_candidates.pop(0)
                origin_path = Path(edited_origin["path"])
                if origin_path.is_absolute():
                    markdown_origin = origin_path.as_posix()
                else:
                    try:
                        markdown_origin = Path(os.path.relpath(edited_origin["resolved"], target.parent)).as_posix()
                    except ValueError:
                        markdown_origin = Path(edited_origin["resolved"]).as_posix()
                origin = json.dumps({"path": markdown_origin, "name": edited_origin["name"]}, ensure_ascii=False)
                blocks.append({"t": "RawBlock", "c": ["html", f"<!-- pptfig-edited {origin} -->"]})
                image_node["c"][2][1] = ""
                blocks.append(block)
                continue
            candidates = marked_images.get(relation, [])
            if not candidates:
                blocks.append(block)
                continue
            metadata, current_markup, word_pixels = candidates.pop(0)
            ppt_path = Path(metadata["resolved"])
            name = metadata["name"]
            unchanged = current_markup == metadata["markup"] and ppt_path.is_file()
            if unchanged:
                try:
                    if ppt_path not in figure_indices:
                        snapshot = media_root / f"deck_{len(deck_snapshots) + 1:06d}.pptx"
                        shutil.copyfile(ppt_path, snapshot)
                        deck_snapshots[ppt_path] = snapshot
                        figure_indices[ppt_path] = read_ppt_figures(snapshot)
                    if name not in figure_indices[ppt_path]:
                        unchanged = False
                    else:
                        key = (ppt_path, name)
                        if key not in rendered_figures:
                            rendered = media_root / f"ppt_compare_{len(rendered_figures) + 1:06d}.png"
                            export_ppt_figures([(deck_snapshots[ppt_path], figure_indices[ppt_path][name][0], rendered)], TEMP_ROOT)
                            rendered_figures[key] = rendered.read_bytes()
                        unchanged = word_pixels == rendered_figures[key]
                except (ValueError, RuntimeError, OSError):
                    unchanged = False
            raw_path = metadata["path"]
            if Path(raw_path).is_absolute():
                markdown_path = raw_path.replace("\\", "/")
            else:
                try:
                    markdown_path = Path(os.path.relpath(ppt_path, target.parent)).as_posix()
                except ValueError:
                    markdown_path = ppt_path.as_posix()
            if unchanged:
                token = f"MDPPTRECOVERED{len(recovered) + 1:06d}"
                recovered[token] = f'{{{{pptfig:"{markdown_path}"|{name}}}}}'
                blocks.append({"t": "Para", "c": [{"t": "Str", "c": token}]})
            else:
                image_node["c"][2][1] = ""
                origin = json.dumps({"path": markdown_path, "name": name}, ensure_ascii=False)
                blocks.append({"t": "RawBlock", "c": ["html", f"<!-- pptfig-edited {origin} -->"]})
                blocks.append(block)
                print(f"提示：Word 图片或 PPT 来源已有变化，已保留 Word 图片并脱离图集引用：{ppt_path} | {name}", file=sys.stderr)
        document["blocks"] = blocks
        images = _export_images(document, media_root, target)
        markdown = run_pandoc(
            pandoc, json.dumps(document, ensure_ascii=False), "json", PANDOC_OUTPUT,
            ["--wrap=none"],
        )
        markdown = re.sub(r"<table\b[^>]*>", "<table>", markdown, flags=re.I)
        markdown = re.sub(r"<colgroup>.*?</colgroup>\s*", "", markdown, flags=re.I | re.S)
        for token, literal in {**citations, **images, **recovered}.items():
            markdown = markdown.replace(token, literal)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(_normalize_markdown_spacing(markdown), encoding="utf-8")
