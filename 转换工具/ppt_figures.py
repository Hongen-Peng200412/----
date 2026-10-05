"""读取 PPT 图名与备注, 并将指定幻灯片导出为裁去页边留白的成品图。"""

from __future__ import annotations

import copy
import json
import locale
import posixpath
import re
import subprocess
import sys
import tempfile
from pathlib import Path, PurePosixPath
from xml.etree import ElementTree
from zipfile import ZipFile


DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
RELATIONSHIP_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
FIGURE_NAME = re.compile(r"@@([^\s|{}]+)\Z")
EXPORT_SCRIPT = Path(__file__).with_name("export_slide.ps1")


def read_ppt_figures(ppt_path: Path) -> dict[str, tuple[int, str]]:
    """返回图名到 (从 1 开始的页码, 原样备注文字) 的映射; 其他批注不参与定位。"""
    from pptx import Presentation

    presentation = Presentation(ppt_path)
    figures: dict[str, tuple[int, str]] = {}
    with ZipFile(ppt_path) as package:
        names = set(package.namelist())
        for page, slide in enumerate(presentation.slides, 1):
            slide_part = slide.part.partname.lstrip("/")
            parent = PurePosixPath(slide_part).parent
            relationships = str(parent / "_rels" / (PurePosixPath(slide_part).name + ".rels"))
            if relationships not in names:
                continue
            root = ElementTree.fromstring(package.read(relationships))
            slide_names = []
            for relation in root.findall(f"{{{RELATIONSHIP_NS}}}Relationship"):
                if not relation.get("Type", "").endswith("/comments"):
                    continue
                comment_part = posixpath.normpath(posixpath.join(str(parent), relation.get("Target", ""))).lstrip("/")
                if comment_part not in names:
                    raise ValueError(f"第 {page} 页批注部件缺失：{comment_part}")
                comments = ElementTree.fromstring(package.read(comment_part))
                for comment in comments:
                    content = "".join(node.text or "" for node in comment.iter(f"{{{DRAWING_NS}}}t")).strip()
                    if not content:
                        content = "".join(node.text or "" for node in comment.iter() if node.tag.rsplit("}", 1)[-1] == "text").strip()
                    matched = FIGURE_NAME.fullmatch(content)
                    if matched:
                        slide_names.append(matched.group(1))
            if len(slide_names) > 1:
                raise ValueError(f"第 {page} 页有多个 @@ 图名批注：{slide_names}")
            if not slide_names:
                continue
            figure_name = slide_names[0]
            if figure_name in figures:
                raise ValueError(f"PPT 图名重复：{figure_name}，第 {figures[figure_name][0]} 与 {page} 页")
            # 保留末尾换行, 使仅含 ENGLISH: 分隔行的空英文区也可被识别; 选择语言后再去除首尾空白.
            note = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
            figures[figure_name] = (page, note)
    return figures


def export_ppt_figures(requests: list[tuple[Path, int, Path]], temp_root: Path) -> None:
    """一次调用 PowerPoint 导出多份 PPT 的指定页, 裁边后分别写出 300 DPI PNG。"""
    from PIL import Image, ImageChops

    if not requests:
        return
    temp_root.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="ppt_render_", dir=temp_root) as directory:
        decks: dict[Path, list[dict[str, str | int]]] = {}
        outputs = []
        for index, (ppt_path, page, output) in enumerate(requests, 1):
            raw = Path(directory) / f"full_slide_{index:06d}.png"
            decks.setdefault(ppt_path, []).append({"page": page, "output": str(raw)})
            outputs.append((raw, output, page))
        manifest = Path(directory) / "slides.json"
        manifest.write_text(
            json.dumps([{"path": str(path), "slides": slides} for path, slides in decks.items()], ensure_ascii=False),
            encoding="utf-8",
        )
        command = [
            "powershell.exe", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
            "-File", str(EXPORT_SCRIPT), "-ManifestPath", str(manifest), "-WidthPx", "3600",
        ]
        for attempt in range(3):
            result = subprocess.run(
                command, capture_output=True, check=False, text=True,
                encoding=locale.getpreferredencoding(False), errors="replace",
            )
            message = "\n".join(part.strip() for part in (result.stderr, result.stdout) if part.strip())
            if result.returncode == 0 and all(raw.is_file() and raw.stat().st_size > 0 for raw, _, _ in outputs):
                break
            if attempt == 2 or "COMException" not in message:
                raise RuntimeError(f"PowerPoint 导图失败：{message or '没有生成全部有效 PNG'}")
        if message:
            print(message, file=sys.stderr)
        for raw, output, page in outputs:
            with Image.open(raw) as original:
                canvas = original.convert("RGB")
                difference = ImageChops.difference(canvas, Image.new("RGB", canvas.size, "white"))
                content = difference.convert("L").point(lambda value: 255 if value > 12 else 0)
                bounds = content.getbbox()
                if bounds is None:
                    raise ValueError(f"PPT 第 {page} 页没有可见图形")
                margin = max(12, round(min(canvas.size) * 0.012))
                left, top, right, bottom = bounds
                cropped = canvas.crop((max(0, left - margin), max(0, top - margin),
                                       min(canvas.width, right + margin), min(canvas.height, bottom + margin)))
                output.parent.mkdir(parents=True, exist_ok=True)
                cropped.save(output, format="PNG", dpi=(300, 300))


def picture_markup(inline) -> str:
    """序列化 Word 图片的视觉属性, 忽略可变的关系编号用于识别直接编辑。"""
    from docx.oxml.ns import qn
    from lxml import etree

    picture = copy.deepcopy(inline.xpath(".//pic:pic")[0])
    for image in picture.xpath(".//a:blip"):
        image.attrib.pop(qn("r:embed"), None)
    extent = inline.xpath("./wp:extent")[0]
    return f"{extent.get('cx')}x{extent.get('cy')}|" + etree.tostring(picture, method="c14n").decode("utf-8")
