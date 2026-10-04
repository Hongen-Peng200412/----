"""用可再生样本核对表格、PPT 图集、引用、公式和 Word 图片编辑的往返行为。"""

from __future__ import annotations

import errno
import json
import shutil
import struct
import subprocess
import sys
import tempfile
import zlib
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.dont_write_bytecode = True

from manuscript_conversion import (
    TEMP_ROOT,
    ensure_dependencies,
    markdown_to_word,
    word_to_markdown,
)
from 转换工具.ppt_figures import export_ppt_figures, read_ppt_figures


def make_captioned_ppt(source: Path, target: Path, caption: str, name_suffix: str = "") -> None:
    """仅在临时 PPT 副本的第 1 页备注中填写图注, 保留真实的现代批注。"""
    from lxml import etree

    p_ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    a_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    with ZipFile(source) as original, ZipFile(target, "w") as updated:
        for item in original.infolist():
            payload = original.read(item.filename)
            if item.filename == "ppt/notesSlides/notesSlide1.xml":
                root = etree.fromstring(payload)
                bodies = root.xpath(
                    ".//p:sp[p:nvSpPr/p:nvPr/p:ph[@type='body']]/p:txBody",
                    namespaces={"p": p_ns},
                )
                assert len(bodies) == 1
                paragraph = bodies[0].find(f"{{{a_ns}}}p")
                run = etree.Element(f"{{{a_ns}}}r")
                etree.SubElement(run, f"{{{a_ns}}}t").text = caption
                paragraph.insert(0, run)
                payload = etree.tostring(root, encoding="UTF-8", xml_declaration=True)
            if name_suffix and item.filename.startswith("ppt/comments/") and item.filename.endswith(".xml"):
                root = etree.fromstring(payload)
                for node in root.iter():
                    if node.text and node.text.startswith("@@"):
                        node.text += name_suffix
                payload = etree.tostring(root, encoding="UTF-8", xml_declaration=True)
            updated.writestr(item, payload)


def test_ppt_figures(root: Path) -> None:
    """用含真实现代批注的图集副本核对多 PPT、图注、回转与 Word 改图。"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    original = ROOT / "画图" / "总览图.pptx"
    if not original.is_file():
        raise FileNotFoundError(f"测试图集不存在：{original}")
    assert "fig-overview" in read_ppt_figures(original)
    first = root / "primary.pptx"
    second = root / "secondary.pptx"
    caption = "总流程示意。"
    make_captioned_ppt(original, first, caption)
    make_captioned_ppt(original, second, caption, "-secondary")
    assert read_ppt_figures(first)["fig-overview"][1] == caption
    source = root / "ppt_article.md"
    source.write_text(
        "# 多图集测试\n\n前文。{{figref:fig-overview-secondary|a–c}}；{{figref:fig-overview|a,b}}。\n"
        '{{pptfig:"primary.pptx"|fig-overview}}\n后文。\n\n'
        f'{{{{pptfig:"{second.as_posix()}"|fig-overview-secondary}}}}\n'
        '组合引用：{{figref:fig-overview,fig-overview-secondary}}。\n'
        '示例代码：`{{figref:not-a-figure}}`。\n',
        encoding="utf-8",
    )
    for profile in ("nature", "operation"):
        reference = ROOT / "论文草稿.operation.docx" if profile == "operation" else None
        first_word = root / f"ppt_{profile}.docx"
        first_markdown = (root / "nested" / "ppt_nature.md") if profile == "nature" else root / "ppt_operation.md"
        second_word = root / f"ppt_{profile}_second.docx"
        second_markdown = first_markdown.with_name(f"ppt_{profile}_second.md")
        with patch("manuscript_conversion.export_ppt_figures", wraps=export_ppt_figures) as exports:
            markdown_to_word(source, first_word, profile=profile, reference=reference)
        assert exports.call_count == 1
        export_requests = exports.call_args.args[0]
        assert len(export_requests) == 2
        assert all(deck.parent.parent == TEMP_ROOT and deck.name.startswith("deck_") for deck, _, _ in export_requests)
        word = Document(first_word)
        assert len(word.inline_shapes) == 2
        assert word.styles["Image Caption"].font.italic is False
        section = word.sections[0]
        text_width = section.page_width - section.left_margin - section.right_margin
        assert abs(word.inline_shapes[0].width - text_width) < 20000
        assert all(
            paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
            for paragraph in word.paragraphs
            if paragraph._p.xpath(".//w:drawing")
        )
        if profile == "operation":
            assert all(
                word.styles[f"Heading {level}"].element.pPr.numPr is None
                for level in range(1, 7)
            )
            captions = [paragraph for paragraph in word.paragraphs
                        if paragraph.text in (f"Fig. 1 | {caption}", f"Fig. 2 | {caption}")]
            assert len(captions) == 2
            assert all(paragraph.paragraph_format.line_spacing == 1.0 for paragraph in captions)
            assert all(
                run.font.name == "华文仿宋" and run.font.size.pt == 9
                for paragraph in captions for run in paragraph.runs
            )
        assert sum(caption in paragraph.text for paragraph in word.paragraphs) == 2
        visible = "\n".join(paragraph.text for paragraph in word.paragraphs)
        assert "Fig. 2a–c；Fig. 1a, b" in visible
        assert "Figs. 1–2" in visible
        assert "{{figref:not-a-figure}}" in visible
        with patch("manuscript_conversion.export_ppt_figures", wraps=export_ppt_figures) as comparisons:
            word_to_markdown(first_word, first_markdown)
        assert comparisons.call_count == 2
        assert all(
            request[0][0].parent.parent == TEMP_ROOT and request[0][0].name.startswith("deck_")
            for request in (call.args[0] for call in comparisons.call_args_list)
        )
        recovered = first_markdown.read_text(encoding="utf-8")
        relative_ppt = "../primary.pptx" if profile == "nature" else "primary.pptx"
        assert f'{{{{pptfig:"{relative_ppt}"|fig-overview}}}}' in recovered
        assert f'{{{{pptfig:"{second.as_posix()}"|fig-overview-secondary}}}}' in recovered
        assert "{{figref:fig-overview-secondary|a–c}}" in recovered
        assert "{{figref:fig-overview|a,b}}" in recovered
        assert "{{figref:fig-overview,fig-overview-secondary}}" in recovered
        markdown_to_word(first_markdown, second_word, profile=profile, reference=reference)
        word_to_markdown(second_word, second_markdown)
        assert first_markdown.read_bytes() == second_markdown.read_bytes()
        if profile == "nature":
            word.inline_shapes[0].width = int(word.inline_shapes[0].width * 0.8)
            edited_word = root / "ppt_edited.docx"
            edited_markdown = root / "ppt_edited.md"
            word.save(edited_word)
            word_to_markdown(edited_word, edited_markdown)
            changed = edited_markdown.read_text(encoding="utf-8")
            assert 'pptfig-edited' in changed and '"name": "fig-overview"' in changed
            assert '&&"' in changed and '{{pptfig:"../primary.pptx"|fig-overview}}' not in changed
            edited_second_word = root / "ppt_edited_second.docx"
            edited_second_markdown = root / "ppt_edited_second.md"
            markdown_to_word(edited_markdown, edited_second_word)
            word_to_markdown(edited_second_word, edited_second_markdown)
            assert edited_markdown.read_bytes() == edited_second_markdown.read_bytes()
            edited_reordered = root / "edited_reordered.md"
            second_directive = f'{{{{pptfig:"{second.as_posix()}"|fig-overview-secondary}}}}'
            edited_reordered.write_text(second_directive + "\n\n" + changed.replace(second_directive, ""), encoding="utf-8")
            edited_reordered_word = root / "edited_reordered.docx"
            markdown_to_word(edited_reordered, edited_reordered_word)
            edited_visible = "\n".join(paragraph.text for paragraph in Document(edited_reordered_word).paragraphs)
            assert "Fig. 1a–c；Fig. 2a, b" in edited_visible
            assert f"Fig. 2 | {caption}" in edited_visible
            replaced = Document(first_word)
            from docx.oxml.ns import qn
            replacement = root / "replacement.png"
            make_png(replacement, (0, 0, 255))
            relation = replaced.inline_shapes[0]._inline.xpath(".//a:blip")[0].get(qn("r:embed"))
            replaced.part.related_parts[relation]._blob = replacement.read_bytes()
            replaced_word = root / "ppt_replaced.docx"
            replaced_markdown = root / "ppt_replaced.md"
            replaced.save(replaced_word)
            word_to_markdown(replaced_word, replaced_markdown)
            assert 'pptfig-edited' in replaced_markdown.read_text(encoding="utf-8")
    reordered = root / "reordered.md"
    reordered.write_text(
        '{{figref:fig-overview-secondary|a}}；{{figref:fig-overview}}。\n\n'
        '{{pptfig:"secondary.pptx"|fig-overview-secondary}}\n\n'
        '{{pptfig:"primary.pptx"|fig-overview}}\n', encoding="utf-8",
    )
    reordered_word = root / "reordered.docx"
    markdown_to_word(reordered, reordered_word)
    visible = "\n".join(paragraph.text for paragraph in Document(reordered_word).paragraphs)
    assert "Fig. 1a；Fig. 2" in visible
    assert f"Fig. 1 | {caption}" in visible and f"Fig. 2 | {caption}" in visible
    conflict = root / "conflict.pptx"
    shutil.copyfile(first, conflict)
    invalid_sources = (
        ('{{pptfig:"primary.pptx"|fig-overview}}\n{{pptfig:"conflict.pptx"|fig-overview_v2}}', "跨 PPT 图名重复"),
        ('{{pptfig:"primary.pptx"|fig-overview}}\n{{pptfig:"primary.pptx"|fig-overview}}', "重复插图"),
        ('{{figref:missing}}\n{{pptfig:"primary.pptx"|fig-overview}}', "没有插图定义"),
        ('{{figref:fig-overview|A}}\n{{pptfig:"primary.pptx"|fig-overview}}', "引用格式无效"),
    )
    for index, (text, expected_error) in enumerate(invalid_sources):
        invalid = root / f"invalid_reference_{index}.md"
        invalid.write_text(text, encoding="utf-8")
        with patch("manuscript_conversion.export_ppt_figures") as exports:
            try:
                markdown_to_word(invalid, root / f"invalid_reference_{index}.docx")
            except ValueError as error:
                assert expected_error in str(error), str(error)
            else:
                raise AssertionError(f"未拒绝无效图引用：{text}")
            exports.assert_not_called()
    missing = root / "missing.md"
    missing.write_text('{{pptfig:"primary.pptx"|not-found}}\n', encoding="utf-8")
    try:
        markdown_to_word(missing, root / "missing.docx")
    except ValueError as error:
        assert "图名不存在" in str(error)
    else:
        raise AssertionError("不存在的 PPT 图名没有报错")
    duplicate = root / "duplicate.pptx"
    with ZipFile(first) as original_zip, ZipFile(duplicate, "w") as changed_zip:
        for item in original_zip.infolist():
            payload = original_zip.read(item.filename)
            if item.filename.startswith("ppt/comments/") and item.filename.endswith(".xml"):
                from lxml import etree
                root_xml = etree.fromstring(payload)
                marker = next(node for node in root_xml if b"@@fig-overview" in etree.tostring(node))
                root_xml.append(etree.fromstring(etree.tostring(marker)))
                payload = etree.tostring(root_xml, encoding="UTF-8", xml_declaration=True)
            changed_zip.writestr(item, payload)
    try:
        read_ppt_figures(duplicate)
    except ValueError as error:
        assert "多个 @@ 图名" in str(error)
    else:
        raise AssertionError("重复的 PPT 图名没有报错")


def make_png(path: Path, rgb: tuple[int, int, int]) -> None:
    """生成一张 12×12 的测试图片，无需额外图像库。"""
    def chunk(name: bytes, payload: bytes) -> bytes:
        content = name + payload
        return struct.pack(">I", len(payload)) + content + struct.pack(">I", zlib.crc32(content))

    row = b"\x00" + bytes(rgb) * 12
    content = b"\x89PNG\r\n\x1a\n"
    content += chunk(b"IHDR", struct.pack(">IIBBBBB", 12, 12, 8, 2, 0, 0, 0))
    content += chunk(b"IDAT", zlib.compress(row * 12))
    content += chunk(b"IEND", b"")
    path.write_bytes(content)


def main() -> None:
    ensure_dependencies()
    from docx import Document
    from openpyxl import Workbook

    TEMP_ROOT.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="manuscript_roundtrip_", dir=TEMP_ROOT) as temporary:
        root = Path(temporary)
        image = root / "figure.png"
        added_image = root / "added.png"
        make_png(image, (255, 0, 0))
        make_png(added_image, (0, 0, 255))
        export_attempts = []

        def intermittent_export(command, **_kwargs):
            export_attempts.append(command)
            if len(export_attempts) == 1:
                return subprocess.CompletedProcess(command, 1, "", "Slide.Export : Object does not exist. COMException")
            manifest = Path(command[command.index("-ManifestPath") + 1])
            decks = json.loads(manifest.read_text(encoding="utf-8"))
            assert len(decks) == 2 and sum(len(deck["slides"]) for deck in decks) == 3
            for deck in decks:
                for slide in deck["slides"]:
                    shutil.copyfile(image, Path(slide["output"]))
            return subprocess.CompletedProcess(command, 0, "", "")

        retried_image = root / "retried.png"
        second_image = root / "second.png"
        third_image = root / "third.png"
        with patch("转换工具.ppt_figures.subprocess.run", side_effect=intermittent_export):
            export_ppt_figures([
                (root / "unused.pptx", 7, retried_image),
                (root / "unused.pptx", 8, second_image),
                (root / "other.pptx", 1, third_image),
            ], root)
        assert len(export_attempts) == 2 and all(path.is_file() for path in (retried_image, second_image, third_image))
        permanent_error = subprocess.CompletedProcess([], 1, "", "Slide.Export : invalid slide")
        with patch("转换工具.ppt_figures.subprocess.run", return_value=permanent_error) as failed_export:
            try:
                export_ppt_figures([(root / "unused.pptx", 99, root / "invalid.png")], root)
            except RuntimeError as error:
                assert "invalid slide" in str(error)
            else:
                raise AssertionError("永久性导图错误未报告")
            assert failed_export.call_count == 1
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Sheet1"
        worksheet.append(["Group", "A", "B"])
        worksheet.append(["One", 10, 20])
        worksheet.append(["", 30, 40])
        worksheet.merge_cells("A2:A3")
        workbook.save(root / "table.xlsx")
        (root / "references.bib").write_text(
            "@article{smith2020, author={Smith, Jane}, title={Example paper}, "
            "journal={Nature Methods}, year={2020}, volume={17}, pages={1--2}}\n",
            encoding="utf-8",
        )
        directive = json.dumps({
            "path": (root / "table.xlsx").as_posix(), "sheet": "Sheet1",
            "range": "A1:C3", "caption": "XLSX table",
        }, ensure_ascii=False)
        source = root / "article.md"
        source.write_text(
            "# Example\n\n第一段。\n第二段 $E=mc^2$ [@smith2020]。\n\n"
            "## 二级\n\n### 三级\n\n#### 四级\n\n##### 五级\n\n###### 六级\n\n"
            "$$f(x)=x^2.$$\n\n"
            "| Metric | Value |\n|---|---:|\n| Score | 0.81 |\n\n"
            "<table><caption>Merged table</caption><tr><th colspan=\"2\">Info</th>"
            "<th>Value</th></tr><tr><td rowspan=\"2\">A</td><td>X</td><td>1</td></tr>"
            "<tr><td>Y</td><td>2</td></tr></table>\n\n"
            f"<!-- xlsx-table {directive} -->\n\n"
            f'![Figure](&&"{image.as_posix()}")\n',
            encoding="utf-8",
        )
        first_word = root / "first.docx"
        first_markdown = root / "first.md"
        second_word = root / "second.docx"
        second_markdown = root / "second.md"
        markdown_to_word(source, first_word, None)
        original_word = first_word.read_bytes()
        pending_before = set(TEMP_ROOT.glob("pending_manuscript_figures_*.docx"))
        with patch("manuscript_conversion.os.replace", side_effect=OSError(errno.EINVAL, "target busy")):
            try:
                markdown_to_word(source, first_word, None)
            except RuntimeError as error:
                assert "完整 Word 已保留" in str(error)
            else:
                raise AssertionError("目标写入失败未报告")
        pending_after = set(TEMP_ROOT.glob("pending_manuscript_figures_*.docx")) - pending_before
        assert len(pending_after) == 1 and first_word.read_bytes() == original_word
        pending_after.pop().unlink()
        word_to_markdown(first_word, first_markdown)
        markdown_to_word(first_markdown, second_word, None)
        word_to_markdown(second_word, second_markdown)
        assert first_markdown.read_bytes() == second_markdown.read_bytes()
        result = first_markdown.read_text(encoding="utf-8")
        assert "| Score" in result
        assert 'colspan="2"' in result and 'rowspan="2"' in result
        assert "[@smith2020]" in result and '&&"' in result
        document = Document(first_word)
        body = [p.text for p in document.paragraphs if p.style.name in ("Body Text", "First Paragraph")]
        assert "第一段。" in body and any("第二段" in item for item in body)
        document.tables[0].cell(1, 1).text = "0.99"
        document.add_paragraph().add_run().add_picture(str(added_image))
        edited_word = root / "edited.docx"
        document.save(edited_word)
        edited_markdown = root / "edited.md"
        word_to_markdown(edited_word, edited_markdown)
        edited = edited_markdown.read_text(encoding="utf-8")
        assert "0.99" in edited and edited.count('&&"') == 2
        if (ROOT / "论文草稿.operation.docx").is_file():
            operation_word = root / "operation.docx"
            markdown_to_word(source, operation_word, profile="operation", reference=ROOT / "论文草稿.operation.docx")
            operation = Document(operation_word)
            levels = {paragraph.style.name for paragraph in operation.paragraphs}
            assert all(f"Heading {level}" in levels for level in range(1, 7))
            assert operation.styles["Heading 4"].font.color.rgb == (0, 0, 0)
            table_captions = [paragraph for paragraph in operation.paragraphs
                              if paragraph._p.pPr is not None and paragraph._p.pPr.pStyle is not None
                              and paragraph._p.pPr.pStyle.val == "TableCaption"]
            assert len(table_captions) == 2
            assert all(paragraph.paragraph_format.line_spacing == 1.0 for paragraph in table_captions)
            assert all(
                run.font.name == "华文仿宋" and run.font.size.pt == 9
                for paragraph in table_captions for run in paragraph.runs
            )
            operation_md = root / "operation.md"
            word_to_markdown(operation_word, operation_md)
            second_operation_word = root / "operation_second.docx"
            markdown_to_word(operation_md, second_operation_word, profile="operation", reference=ROOT / "论文草稿.operation.docx")
            second_operation_md = root / "operation_second.md"
            word_to_markdown(second_operation_word, second_operation_md)
            assert operation_md.read_bytes() == second_operation_md.read_bytes()
        test_ppt_figures(root)
        print("通过：两种版式、多 PPT 图集、图注、三类表格、公式、引用和 Word 内图片编辑。")


if __name__ == "__main__":
    main()
