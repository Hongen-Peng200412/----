"""通过正式 Word 流程验证双语 PPT 图注, 固定图片以隔离 Office 桌面状态。"""

from __future__ import annotations

import hashlib
import sys
import tempfile
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.dont_write_bytecode = True

from manuscript_conversion import TEMP_ROOT, ensure_dependencies, markdown_to_word, word_to_markdown
from 转换工具.ppt_figures import read_ppt_figures
from 转换工具.test_roundtrip import make_captioned_ppt


def main() -> None:
    """核对真实备注读取、语言选择、两种版式、标题格式与英文往返。"""
    ensure_dependencies()
    from docx import Document
    from PIL import Image

    chinese = "中文标题。中文说明。"
    english = "English title. English description."
    panel_chinese = "面板标题。\na-c，短横线范围。\nd–f、连接号范围。\n  g - i，含空格范围。\nj – l，含空格连接号范围。\nm，单个面板。\n**n–p**，已有加粗。\n比较 a-c，普通说明。\na-c 不是段首面板标识。"
    panel_english = "Panel title.\na-c, Hyphen range.\nd–f, En dash range.\n  g - i, Spaced hyphen range.\nj – l, Spaced en dash range.\nm, Single panel.\n**n–p**, Existing bold.\nCompare a-c, ordinary prose.\na-c bonds remain unchanged."
    cases = [
        (chinese + "\nENGLISH:\n" + english, chinese, english),
        (chinese + "\nENGLISH：\n" + english, chinese, english),
        (chinese + "\r\nENGLISH:\r\n" + english, chinese, english),
        (chinese + "\rENGLISH：\r" + english, chinese, english),
        (chinese + "\n \tENGLISH: \t\n" + english, chinese, english),
        (chinese, chinese, chinese),
        (english, english, english),
        (chinese + " ENGLISH:\n" + english, chinese + " ENGLISH:\n" + english,
         chinese + " ENGLISH:\n" + english),
        (chinese + "\nENGLISH:\\n" + english, chinese + "\nENGLISH:\\n" + english,
         chinese + "\nENGLISH:\\n" + english),
        (chinese + "\nENGLISH:", chinese + "\nENGLISH:", chinese + "\nENGLISH:"),
        (chinese + "\nEnglish:\n" + english, chinese + "\nEnglish:\n" + english,
         chinese + "\nEnglish:\n" + english),
        (chinese + "\nENGLISH:\n", chinese, ""),
        ("ENGLISH:\n" + english, "", english),
        (chinese + "\nENGLISH:\n" + english + "\nENGLISH:\nTail.", chinese,
         english + "\nENGLISH:\nTail."),
        ("图 9｜" + chinese + "\nENGLISH:\nFig. 8 | " + english, chinese, english),
        (chinese + "\nENGLISH:\nEnglish title. Native formula $h_j^{\\mathrm{vox}}$ remains editable.",
         chinese, "English title. Native formula  remains editable."),
        (panel_chinese + "\nENGLISH:\n" + panel_english,
         panel_chinese.replace("**", ""), panel_english.replace("**", "")),
    ]
    TEMP_ROOT.mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="english_captions_", dir=TEMP_ROOT) as folder:
        work = Path(folder)
        source = work / "captions.md"
        decks = []
        for index, (notes, _, _) in enumerate(cases):
            deck = work / f"case_{index}.pptx"
            suffix = f"-english-test-{index}"
            make_captioned_ppt(ROOT / "画图" / "总览图.pptx", deck, notes, suffix)
            name = "fig-overview" + suffix
            assert read_ppt_figures(deck)[name][1] == notes, f"备注读取改变了 case {index}"
            decks.append(deck)
        source.write_text(
            "# Caption selection\n\nBefore {{figref:fig-overview-english-test-0|a,b}}.\n\n"
            + "\n\n".join(
                f'{{{{pptfig:"{deck.as_posix()}"|fig-overview-english-test-{index}}}}}'
                for index, deck in enumerate(decks)
            ), encoding="utf-8",
        )
        original_hashes = [hashlib.sha256(deck.read_bytes()).hexdigest() for deck in decks]

        # 图片固定后仍执行真实 Pandoc、DOCX 格式化及回转, 可检查图注与编号的组合行为.
        with patch("manuscript_conversion.export_ppt_figures") as exports:
            exports.side_effect = lambda requests, _: [
                Image.new("RGB", (320, 120), "white").save(target)
                for _, _, target in requests
            ]
            for profile in ("nature", "operation"):
                reference = ROOT / "论文草稿.operation.docx" if profile == "operation" else None
                for use_english in (False, True):
                    target = work / f"{profile}_{use_english}.docx"
                    markdown_to_word(source, target, profile=profile, reference=reference,
                                     english=use_english)
                    word = Document(target)
                    captions = [paragraph for paragraph in word.paragraphs
                                if paragraph.text.startswith("Fig. ")]
                    # 字面的反斜杠由 Pandoc 按 Markdown/TeX 解释; 此例只检查语言标记未被当作分区.
                    expected = [f"Fig. {index + 1} | {case[2 if use_english else 1]}"
                                for index, case in enumerate(cases)
                                if index != 8 and case[2 if use_english else 1]]
                    literal_caption = next(p.text for p in captions if p.text.startswith("Fig. 9 | "))
                    assert "ENGLISH:" in literal_caption
                    actual = [paragraph.text for paragraph in captions
                              if not paragraph.text.startswith("Fig. 9 | ")]
                    assert [" ".join(text.split()) for text in actual] == [
                        " ".join(text.split()) for text in expected
                    ], (profile, use_english, actual, expected)
                    assert len(word.inline_shapes) == len(cases)
                    assert "Before Fig. 1a, b." in "\n".join(p.text for p in word.paragraphs)
                    first = captions[0]
                    title = "English title." if use_english else "中文标题。"
                    assert "".join(run.text for run in first.runs if run.bold) == "Fig. 1 | " + title
                    panels = next(p for p in captions if p.text.startswith(f"Fig. {len(cases)} | "))
                    panel_title = "Panel title." if use_english else "面板标题。"
                    assert "".join(run.text for run in panels.runs if run.bold) == (
                        f"Fig. {len(cases)} | {panel_title}" + "a-cd–fg - ij – lmn–p"
                    ), (profile, use_english, [(run.text, run.bold) for run in panels.runs])
                    assert all(run.italic is not True for paragraph in captions for run in paragraph.runs)
                    if profile == "operation":
                        assert all(p.paragraph_format.line_spacing == 1.0 for p in captions)
                        assert all(run.font.name == "华文仿宋" and run.font.size.pt == 9
                                   for p in captions for run in p.runs)
                    if use_english:
                        assert word._element.xpath(".//m:oMath"), "英文图注公式未生成原生公式"
                        returned = work / f"{profile}_returned.md"
                        word_to_markdown(target, returned)
                        assert returned.read_text(encoding="utf-8").count("{{pptfig:") == len(cases)
                        rebuilt = work / f"{profile}_rebuilt.docx"
                        markdown_to_word(returned, rebuilt, profile=profile, reference=reference,
                                         english=True)
                        assert [p.text for p in Document(rebuilt).paragraphs if p.text.startswith("Fig. ")] == [
                            p.text for p in captions
                        ]
                        assert ["".join(run.text for run in p.runs if run.bold)
                                for p in Document(rebuilt).paragraphs if p.text.startswith("Fig. ")] == [
                            "".join(run.text for run in p.runs if run.bold) for p in captions
                        ]
        assert [hashlib.sha256(deck.read_bytes()).hexdigest() for deck in decks] == original_hashes
        print(f"通过: {len(cases)} 种备注边界 × 2 种语言 × 2 种版式; 原生公式、英文回转、图号、格式和 PPT 保全.")


if __name__ == "__main__":
    main()
