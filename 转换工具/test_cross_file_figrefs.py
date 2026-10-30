"""在temp合成稿件验证跨文件图引用, 实际使用Pandoc/DOCX, 不打开Office或修改论文图集."""

from pathlib import Path
import json
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.dont_write_bytecode = True
import manuscript_conversion as converter


class CrossFileFigureReferenceTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix='cross_file_figrefs_test_', dir=ROOT / 'temp')
        self.root = Path(self.temporary.name)
        self.foreign = self.root / '正文 main.md'
        self.foreign.write_text('{{pptfig:"unused.pptx"|shared}}\n\n{{pptfig:"unused.pptx"|remote}}\n', encoding='utf-8')
        self.local = {'shared': 7}
        self.pandoc = converter.ensure_dependencies()

    def tearDown(self):
        self.temporary.cleanup()

    def encode(self, text):
        document = converter._as_json(self.pandoc, text, converter.PANDOC_INPUT)
        converter._encode_references(document, [], self.local, self.root, {})
        return document

    def test_separate_numbering_spaces_panels_ranges_and_code_examples(self):
        text = '本稿{{figref:shared}}；外稿{{正文 main.md|figref:remote|a,d}}；组合{{正文 main.md|figref:shared,remote}}。\n\n`{{missing.md|figref:missing}}`\n\n```text\n{{missing.md|figref:missing}}\n```'
        document = self.encode(text)
        encoded = json.dumps(document, ensure_ascii=False)
        self.assertIn('Fig. 7', encoded)
        self.assertIn('Fig. 2a, d', encoded)
        self.assertIn('Figs. 1–2', encoded)
        self.assertEqual(encoded.count('{{missing.md|figref:missing}}'), 2)

    def test_reorder_and_edited_figure_definition_without_recursive_compilation(self):
        self.foreign.write_text('{{figref:unresolved}}\n\n<!-- pptfig-edited {"name":"remote","path":"unused.pptx"} -->\n\n{{pptfig:"unused.pptx"|shared}}\n', encoding='utf-8')
        encoded = json.dumps(self.encode('{{正文 main.md|figref:remote}}'), ensure_ascii=False)
        self.assertIn('Fig. 1', encoded)
        self.assertNotIn('unresolved', encoded)

    def test_target_missing_duplicate_and_local_missing_remain_errors(self):
        with self.assertRaises(FileNotFoundError):
            self.encode('{{missing.md|figref:remote}}')
        with self.assertRaisesRegex(ValueError, '没有插图定义'):
            self.encode('{{正文 main.md|figref:absent}}')
        with self.assertRaisesRegex(ValueError, '没有插图定义'):
            self.encode('{{figref:remote}}')
        self.foreign.write_text('{{pptfig:"unused.pptx"|remote}}\n{{pptfig:"unused.pptx"|remote}}\n', encoding='utf-8')
        with self.assertRaisesRegex(ValueError, '重复插图'):
            self.encode('{{正文 main.md|figref:remote}}')

    def test_invalid_panel_and_combined_panel_remain_errors(self):
        for text in ('{{正文main.md|figref:remote|A}}', '{{正文 main.md|figref:shared,remote|a}}'):
            with self.assertRaisesRegex(ValueError, '格式无效|组合引用不能附加面板'):
                self.encode(text)

    def test_one_compile_uses_one_snapshot(self):
        cache = {}
        for expected in (2, 2):
            document = converter._as_json(self.pandoc, '{{正文 main.md|figref:remote}}', converter.PANDOC_INPUT)
            converter._encode_references(document, [], {}, self.root, cache)
            self.assertIn(f'Fig. {expected}', json.dumps(document))
            self.foreign.write_text('{{pptfig:"unused.pptx"|remote}}\n', encoding='utf-8')
        self.assertIn('Fig. 1', json.dumps(self.encode('{{正文 main.md|figref:remote}}')))

    def test_real_docx_roundtrip_and_reference_only_does_not_export_images(self):
        from docx import Document
        text = '本稿{{figref:shared}}；外稿{{正文 main.md|figref:remote|a,d}}。\n\n{{pptfig:"local.pptx"|shared}}\n'
        source = self.root / 'local.md'
        source.write_text(text, encoding='utf-8')
        (self.root / 'local.pptx').write_bytes(b'synthetic fixture, not a manuscript slide')
        # 单色图片仅替代测试夹具的Office导图, 实际编译/样式/回转使用生产流程.
        from PIL import Image
        def export(requests, temp_root):
            for _, _, output in requests:
                output.parent.mkdir(parents=True, exist_ok=True)
                Image.new('RGB', (600, 300), '#246789').save(output, dpi=(300, 300))
        for profile in ('nature', 'operation'):
            with self.subTest(profile=profile):
                word = self.root / 'compiled' / f'{profile}.docx'
                target = self.root / 'nested' / f'{profile}.md'
                reference = ROOT / 'supplementary.operation.docx' if profile == 'operation' else None
                with patch.object(converter, 'read_ppt_figures', return_value={'shared': (1, '固定图注。')}), patch.object(converter, 'export_ppt_figures', side_effect=export) as exports:
                    converter.markdown_to_word(source, word, profile=profile, reference=reference)
                    self.assertEqual(len(exports.call_args.args[0]), 1)
                    actual = Document(word)
                    visible = '\n'.join(p.text for p in actual.paragraphs)
                    self.assertIn('Fig. 1；外稿Fig. 2a, d', visible)
                    self.assertEqual(len(actual.inline_shapes), 1)
                    self.assertIn('Fig. 1 | 固定图注。', visible)
                    converter.word_to_markdown(word, target)
                    recovered = target.read_text(encoding='utf-8')
                    self.assertIn('{{../正文 main.md|figref:remote|a,d}}', recovered)
                    self.assertIn('{{figref:shared}}', recovered)
                    second_word = self.root / f'{profile}-again.docx'
                    second_md = target.with_name(f'{profile}-again.md')
                    converter.markdown_to_word(target, second_word, profile=profile, reference=reference)
                    converter.word_to_markdown(second_word, second_md)
                    self.assertEqual(target.read_bytes(), second_md.read_bytes())
        reference_only = self.root / 'reference-only.md'
        reference_only.write_text('{{正文 main.md|figref:remote}}', encoding='utf-8')
        with patch.object(converter, 'export_ppt_figures') as exports:
            converter.markdown_to_word(reference_only, self.root / 'reference-only.docx')
            exports.assert_called_once_with([], unittest.mock.ANY)
            self.assertEqual(len(Document(self.root / 'reference-only.docx').inline_shapes), 0)


if __name__ == '__main__':
    unittest.main(verbosity=2)
