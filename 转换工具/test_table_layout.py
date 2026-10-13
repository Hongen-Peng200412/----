"""验证阅读版表格不继承正文缩进, 同一比较记录不拆跨页; 所有样本留在论文 temp/."""

import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.dont_write_bytecode = True

from manuscript_conversion import ensure_dependencies, _apply_operation_layout


class OperationTableTest(unittest.TestCase):
    def test_table_cells_and_body_keep_separate_layout(self):
        ensure_dependencies()
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Pt

        with tempfile.TemporaryDirectory(dir=ROOT / "temp", prefix="table_layout_") as folder:
            path = Path(folder) / "table.docx"
            reference = Path(folder) / "reference.docx"
            Document().save(reference)
            document = Document()
            document.add_paragraph("Body remains unchanged")
            table = document.add_table(rows=2, cols=3)
            values = (("Scope", "Model", "PRAUC"), ("Small molecules", "LigandSeek2", "0.572615"))
            for row, entries in zip(table.rows, values):
                for cell, value in zip(row.cells, entries):
                    cell.text = value
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            document.save(path)
            _apply_operation_layout(path, reference, "center", "fit")
            actual = Document(path)
            self.assertEqual(actual.paragraphs[0].text, "Body remains unchanged")
            self.assertEqual(actual.paragraphs[0].paragraph_format.first_line_indent, Pt(21))
            self.assertEqual(actual.paragraphs[0].paragraph_format.line_spacing, 1.5)
            for row, entries in zip(actual.tables[0].rows, values):
                self.assertIsNotNone(row._tr.get_or_add_trPr().find(qn("w:cantSplit")))
                for cell, value in zip(row.cells, entries):
                    self.assertEqual(cell.text, value)
                    format_ = cell.paragraphs[0].paragraph_format
                    self.assertEqual(format_.first_line_indent, Pt(0))
                    self.assertEqual(format_.line_spacing, 1.0)
                    self.assertEqual(format_.space_after, Pt(3))
            self.assertIsNotNone(actual.tables[0].rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")))


if __name__ == "__main__":
    unittest.main()
